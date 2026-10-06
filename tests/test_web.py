"""End-to-end tests of logins and per-user chats, with ConnectWise and Claude mocked out."""

import importlib
import logging

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def web(monkeypatch, tmp_path):
    for name in ("CW_SITE", "CW_COMPANY_ID", "CW_PUBLIC_KEY", "CW_PRIVATE_KEY", "CW_CLIENT_ID"):
        monkeypatch.setenv(name, "x")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    monkeypatch.setenv("CLAUDE_MODEL", "claude-opus-5-5")
    monkeypatch.setenv("CLAUDE_MODELS", "claude-opus-5-5,claude-haiku-4-5")
    monkeypatch.setenv("DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("USERS_FILE", str(tmp_path / "users.txt"))
    (tmp_path / "users.txt").write_text(
        "alice | Alice A | password-a | admin\nbob | Bob B | password-b |\ncarol | Carol C | password-c | uploader\n", encoding="utf-8"
    )
    from dbs_reporting import web as module

    module = importlib.reload(module)
    from dbs_reporting import activity

    monkeypatch.setattr(activity, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(activity, "_logger", None)
    for handler in list(logging.getLogger("dbs_reporting.activity").handlers):
        logging.getLogger("dbs_reporting.activity").removeHandler(handler)
    calls = []

    def fake_respond_stream(history, question, model=None):
        calls.append((list(history), question, model))
        yield {"type": "status", "text": "Looking up the company…"}
        yield {"type": "reset"}
        answer = f"answer to {question}"
        yield {"type": "text", "text": answer[:6]}
        yield {"type": "text", "text": answer[6:]}
        yield {"type": "done", "answer": answer, "history": history + [
            {"role": "user", "content": question},
            {"role": "assistant", "content": [{"type": "text", "text": "a"}]},
        ], "usage": {"requests": 2, "input_tokens": 12000, "output_tokens": 800, "cache_read_tokens": 3000,
                     "cache_write_tokens": 0, "cost_usd": 0.0326, "priced": True}}

    monkeypatch.setattr(module.agent, "respond_stream", fake_respond_stream)
    return module, calls


def login(module, username, password):
    client = TestClient(module.app)
    response = client.post("/api/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    return client


def test_requires_login(web):
    module, _ = web
    client = TestClient(module.app)
    assert client.get("/", follow_redirects=False).headers["location"] == "/login"
    assert client.get("/api/conversations").status_code == 401
    assert client.post("/api/chat", json={"question": "hi"}).status_code == 401
    bad = client.post("/api/login", json={"username": "alice", "password": "wrong"})
    assert bad.status_code == 401


def test_chats_are_saved_and_private(web):
    module, calls = web
    alice = login(module, "alice", "password-a")
    bob = login(module, "bob", "password-b")

    first = alice.post("/api/chat", json={"question": "Issues at Jimmy's Grille?", "model": "claude-haiku-4-5"}).json()
    assert first["answer"] == "answer to Issues at Jimmy's Grille?"
    assert first["title"] == "Issues at Jimmy's Grille?"
    chat_id = first["conversation_id"]

    # Follow-up continues the saved history with the chat's own model.
    alice.post("/api/chat", json={"question": "and last week?", "conversation_id": chat_id})
    history, question, model = calls[-1]
    assert question == "and last week?" and model == "claude-haiku-4-5" and len(history) == 2

    chats = alice.get("/api/conversations").json()
    assert [c["id"] for c in chats] == [chat_id]
    turns = alice.get(f"/api/conversations/{chat_id}").json()["turns"]
    assert [t["role"] for t in turns] == ["user", "assistant", "user", "assistant"]

    # Bob can't see, continue, rename or delete Alice's chat.
    assert bob.get("/api/conversations").json() == []
    assert bob.get(f"/api/conversations/{chat_id}").status_code == 404
    assert bob.post("/api/chat", json={"question": "x", "conversation_id": chat_id}).status_code == 404
    assert bob.patch(f"/api/conversations/{chat_id}", json={"title": "mine"}).status_code == 404
    assert bob.delete(f"/api/conversations/{chat_id}").status_code == 404

    assert alice.patch(f"/api/conversations/{chat_id}", json={"title": "Jimmy's"}).status_code == 200
    assert alice.get("/api/conversations").json()[0]["title"] == "Jimmy's"
    assert alice.delete(f"/api/conversations/{chat_id}").status_code == 200
    assert alice.get("/api/conversations").json() == []


def test_chats_survive_restart(web, monkeypatch):
    module, _ = web
    alice = login(module, "alice", "password-a")
    chat_id = alice.post("/api/chat", json={"question": "hello"}).json()["conversation_id"]

    from dbs_reporting.store import Store

    reopened = Store(module.store.path)
    user = reopened.authenticate("alice", "password-a")
    assert [c["id"] for c in reopened.list_conversations(user["id"])] == [chat_id]
    assert len(reopened.get_conversation(user["id"], chat_id)["history"]) == 2


def test_logout_and_lockout(web):
    module, _ = web
    alice = login(module, "alice", "password-a")
    alice.post("/api/logout")
    assert alice.get("/api/me").status_code == 401

    client = TestClient(module.app)
    for _ in range(5):
        client.post("/api/login", json={"username": "bob", "password": "nope"})
    locked = client.post("/api/login", json={"username": "bob", "password": "password-b"})
    assert locked.status_code == 429


def test_rejects_disabled_model(web):
    module, _ = web
    alice = login(module, "alice", "password-a")
    response = alice.post("/api/chat", json={"question": "hi", "model": "claude-bogus"})
    assert response.status_code == 400


def test_removing_from_users_file_signs_out(web, tmp_path):
    import os

    module, _ = web
    bob = login(module, "bob", "password-b")
    assert bob.get("/api/me").status_code == 200
    path = tmp_path / "users.txt"
    path.write_text("\n".join(l for l in path.read_text().splitlines() if not l.startswith("bob")) + "\n")
    stamp = path.stat().st_mtime + 10
    os.utime(path, (stamp, stamp))
    assert bob.get("/api/me").status_code == 401


def test_model_is_locked_per_chat(web):
    module, calls = web
    alice = login(module, "alice", "password-a")
    first = alice.post("/api/chat", json={"question": "hi", "model": "claude-haiku-4-5"}).json()
    assert first["model"] == "claude-haiku-4-5"
    chat_id = first["conversation_id"]

    # Asking for a different model on an existing chat still uses the chat's own model.
    second = alice.post("/api/chat", json={"question": "again", "conversation_id": chat_id,
                                           "model": "claude-opus-5-5"}).json()
    assert second["conversation_id"] == chat_id and second["model"] == "claude-haiku-4-5"
    assert calls[-1][2] == "claude-haiku-4-5"

    chat = alice.get(f"/api/conversations/{chat_id}").json()
    assert chat["model"] == "claude-haiku-4-5"
    assert [t["model"] for t in chat["turns"] if t["role"] == "assistant"] == ["claude-haiku-4-5"] * 2


def test_disabled_model_falls_back_to_default(web, monkeypatch):
    module, calls = web
    alice = login(module, "alice", "password-a")
    chat_id = alice.post("/api/chat", json={"question": "hi", "model": "claude-haiku-4-5"}).json()["conversation_id"]
    monkeypatch.setattr(module.agent, "models", ["claude-opus-5-5"])
    response = alice.post("/api/chat", json={"question": "again", "conversation_id": chat_id})
    assert response.status_code == 200 and response.json()["model"] == "claude-opus-5-5"


def test_logo(web, tmp_path, monkeypatch):
    module, _ = web
    client = TestClient(module.app)
    monkeypatch.setattr(module, "BRANDING", tmp_path / "branding")
    assert client.get("/logo").status_code == 404
    (tmp_path / "branding").mkdir()
    (tmp_path / "branding" / "logo.png").write_bytes(b"\x89PNG fake")
    response = client.get("/logo")  # public: no sign-in needed
    assert response.status_code == 200 and response.headers["content-type"] == "image/png"


def test_dark_logo(web, tmp_path, monkeypatch):
    module, _ = web
    client = TestClient(module.app)
    branding = tmp_path / "branding"
    branding.mkdir()
    monkeypatch.setattr(module, "BRANDING", branding)
    (branding / "logo.svg").write_text("<svg/>")
    assert client.get("/logo-dark").status_code == 404  # optional; pages fall back to /logo
    (branding / "logo-dark.png").write_bytes(b"\x89PNG fake")
    response = client.get("/logo-dark")
    assert response.status_code == 200 and response.headers["content-type"] == "image/png"
    assert client.get("/logo").headers["content-type"].startswith("image/svg")


def test_logo_is_built_into_the_pages(web, tmp_path, monkeypatch):
    import base64

    module, _ = web
    branding = tmp_path / "branding"
    branding.mkdir()
    monkeypatch.setattr(module, "BRANDING", branding)
    client = TestClient(module.app)
    # No logo: the built-in icon stays.
    assert "<!--brand-mark-->" in client.get("/login").text and "<picture>" not in client.get("/login").text

    (branding / "logo.svg").write_text("<svg>light</svg>")
    page = client.get("/login").text
    light = "data:image/svg+xml;base64," + base64.b64encode(b"<svg>light</svg>").decode()
    assert f'<span class="brand-mark has-logo" aria-hidden="true"><picture><img src="{light}"' in page  # no swap later
    assert "<!--brand-mark-->" not in page and 'media="(prefers-color-scheme: dark)"' not in page

    (branding / "logo-dark.png").write_bytes(b"\x89PNG dark")
    alice = login(module, "alice", "password-a")
    chat = alice.get("/").text
    dark = "data:image/png;base64," + base64.b64encode(b"\x89PNG dark").decode()
    assert f'<source data-dark srcset="{dark}" media="(prefers-color-scheme: dark)">' in chat and light in chat
    assert '<link rel="icon" href="/logo-dark" media="(prefers-color-scheme: dark)">' in chat

    # Large logos are linked rather than embedded in every page.
    (branding / "logo.svg").write_text("<svg>" + "x" * 200_000 + "</svg>")
    assert '<img src="/logo"' in client.get("/login").text


def read_events(response) -> list[dict]:
    import json

    return [json.loads(line) for line in response.text.splitlines() if line.strip()]


def test_stream_endpoint(web):
    module, calls = web
    alice = login(module, "alice", "password-a")
    response = alice.post("/api/chat/stream", json={"question": "Issues at Jimmy's?", "model": "claude-haiku-4-5"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/x-ndjson")
    events = read_events(response)
    assert [e["type"] for e in events] == ["status", "reset", "text", "text", "done"]
    assert "".join(e["text"] for e in events if e["type"] == "text") == "answer to Issues at Jimmy's?"
    done = events[-1]
    assert done["title"] == "Issues at Jimmy's?" and done["model"] == "claude-haiku-4-5"

    # The streamed answer is saved like any other, and follow-ups continue the chat.
    chat = alice.get(f"/api/conversations/{done['conversation_id']}").json()
    assert [t["text"] for t in chat["turns"]] == ["Issues at Jimmy's?", "answer to Issues at Jimmy's?"]
    follow = read_events(alice.post("/api/chat/stream", json={"question": "and last week?",
                                                             "conversation_id": done["conversation_id"]}))
    assert follow[-1]["conversation_id"] == done["conversation_id"]
    assert len(calls[-1][0]) == 2 and calls[-1][2] == "claude-haiku-4-5"


def test_stream_errors(web, monkeypatch):
    module, _ = web
    alice = login(module, "alice", "password-a")
    assert TestClient(module.app).post("/api/chat/stream", json={"question": "hi"}).status_code == 401
    assert alice.post("/api/chat/stream", json={"question": "  "}).status_code == 400
    assert alice.post("/api/chat/stream", json={"question": "x", "conversation_id": "nope"}).status_code == 404

    def broken(history, question, model=None):
        yield {"type": "status", "text": "Looking up the company…"}
        raise RuntimeError("ConnectWise fell over")

    monkeypatch.setattr(module.agent, "respond_stream", broken)
    logged = []
    monkeypatch.setattr(module.log, "exception", lambda msg, *a, **k: logged.append(msg))
    events = read_events(alice.post("/api/chat/stream", json={"question": "hi"}))
    assert events[-1] == {"type": "error", "message": "ConnectWise fell over"}
    assert alice.get("/api/conversations").json() == []  # nothing half-saved
    assert logged == ["Agent error"]  # logged once, not again when shown to the user
    alice.post("/api/chat", json={"question": "hi"})
    assert logged == ["Agent error", "Agent error"]


def test_api_errors_show_the_apis_message(web, monkeypatch):
    import anthropic
    import httpx2

    module, _ = web
    alice = login(module, "alice", "password-a")

    def rejected(history, question, model=None):
        request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
        response = httpx2.Response(400, request=request, json={
            "type": "error", "error": {"type": "invalid_request_error",
                                       "message": "messages.5.content.0: Invalid `signature` in `thinking` block."}})
        raise anthropic.BadRequestError("bad request", response=response, body=response.json())
        yield  # pragma: no cover

    monkeypatch.setattr(module.agent, "respond_stream", rejected)
    events = read_events(alice.post("/api/chat/stream", json={"question": "hi"}))
    assert events[-1]["type"] == "error"
    assert events[-1]["message"].startswith("AI service error (400): messages.5.content.0: Invalid `signature`")


def test_usage_visible_to_admins_only(web):
    module, _ = web
    alice = login(module, "alice", "password-a")  # admin
    bob = login(module, "bob", "password-b")
    assert alice.get("/api/me").json()["is_admin"] is True and bob.get("/api/me").json()["is_admin"] is False

    a = read_events(alice.post("/api/chat/stream", json={"question": "hi"}))[-1]
    b = read_events(bob.post("/api/chat/stream", json={"question": "hi"}))[-1]
    assert a["usage"]["input_tokens"] == 12000 and b["usage"] is None

    a_turns = alice.get(f"/api/conversations/{a['conversation_id']}").json()["turns"]
    b_turns = bob.get(f"/api/conversations/{b['conversation_id']}").json()["turns"]
    assert a_turns[1]["usage"]["cost_usd"] == 0.0326 and b_turns[1]["usage"] is None

    # Both answers are recorded for the usage report, whoever asked.
    from dbs_reporting import usage

    text = usage.report(days=30)
    assert "Answers: 2" in text and "Alice A" in text and "Bob B" in text and "$0.07" in text


def test_activity_log_and_history(web, tmp_path, monkeypatch):
    from dbs_reporting import activity

    module, _ = web
    alice = login(module, "alice", "password-a")
    bob = login(module, "bob", "password-b")
    read_events(alice.post("/api/chat/stream", json={"question": "Printer issues at Jimmy's?"}))
    read_events(bob.post("/api/chat/stream", json={"question": "Hours at Burger Barn?"}))

    def broken(history, question, model=None):
        raise RuntimeError("ConnectWise fell over")
        yield  # pragma: no cover

    monkeypatch.setattr(module.agent, "respond_stream", broken)
    read_events(bob.post("/api/chat/stream", json={"question": "Will this work?"}))

    log = (tmp_path / "logs" / "activity.log").read_text(encoding="utf-8")
    assert "Alice A (alice)" in log and "Printer issues at Jimmy's?" in log
    assert "answer to Printer issues at Jimmy's?" in log
    assert "Usage: 15.0k in (3.0k cached) · 800 out · 2 calls · ≈ $0.033" in log
    assert "Bob B (bob)" in log and "ERROR" in log and "ConnectWise fell over" in log

    everything = activity.history_report(days=7, user=None, search=None, full=False, limit=50)
    assert "Printer issues" in everything and "Burger Barn" in everything and "2 answer(s) shown." in everything
    only_bob = activity.history_report(days=7, user="bob", search=None, full=False, limit=50)
    assert "Burger Barn" in only_bob and "Printer issues" not in only_bob
    searched = activity.history_report(days=7, user=None, search="printer", full=False, limit=50)
    assert "Printer issues" in searched and "Burger Barn" not in searched


def test_tool_calls_are_described():
    from dbs_reporting.activity import tool_calls

    claude = [
        {"role": "user", "content": "q"},
        {"role": "assistant", "content": [
            {"type": "thinking", "thinking": "", "signature": "x"},
            {"type": "tool_use", "id": "t1", "name": "get_ticket_totals", "input": {"days": 30, "group_by": "site"}}]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "t1", "content": "{}"}]},
        {"role": "assistant", "content": [{"type": "text", "text": "done"}]},
    ]
    assert tool_calls(claude) == ["get_ticket_totals(days=30, group_by='site')"]


def test_charts_and_excel_export(web, monkeypatch):
    import io

    from openpyxl import load_workbook

    module, _ = web
    chart_input = {"title": "Tickets by site", "chart_type": "hbar", "labels": ["Downtown", "Airport"],
                   "series": [{"name": "Tickets", "values": [31, 18]}]}

    def with_chart(history, question, model=None):
        answer = "| Site | Tickets |\n|---|---|\n| Downtown | 31 |\n| Airport | 18 |"
        yield {"type": "text", "text": answer}
        yield {"type": "done", "answer": answer, "usage": None, "history": history + [
            {"role": "user", "content": question},
            {"role": "assistant", "content": [{"type": "tool_use", "id": "c1", "name": "create_chart",
                                               "input": chart_input}]},
            {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "c1",
                                          "content": '{"chart_added": true}'}]},
            {"role": "assistant", "content": [{"type": "text", "text": answer}]},
        ]}

    monkeypatch.setattr(module.agent, "respond_stream", with_chart)
    alice = login(module, "alice", "password-a")
    done = read_events(alice.post("/api/chat/stream", json={"question": "sites?"}))[-1]
    assert [c["title"] for c in done["charts"]] == ["Tickets by site"] and done["answer_id"]

    turns = alice.get(f"/api/conversations/{done['conversation_id']}").json()["turns"]
    assert turns[1]["charts"][0]["labels"] == ["Downtown", "Airport"] and turns[1]["id"] == done["answer_id"]

    response = alice.get(f"/api/answers/{done['answer_id']}/export.xlsx")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Table 1", "Chart 1"]

    # Only the person who asked can download it.
    bob = login(module, "bob", "password-b")
    assert bob.get(f"/api/answers/{done['answer_id']}/export.xlsx").status_code == 404
    assert module.app.url_path_for("static", path="vendor/chart.umd.min.js")
    assert bob.get("/static/vendor/chart.umd.min.js").status_code == 200

    # A plain answer (no table or chart) has nothing to download.
    def plain_answer(history, question, model=None):
        yield {"type": "done", "answer": "Just text.", "usage": None,
               "history": history + [{"role": "user", "content": question},
                                     {"role": "assistant", "content": [{"type": "text", "text": "Just text."}]}]}

    monkeypatch.setattr(module.agent, "respond_stream", plain_answer)
    plain = alice.post("/api/chat", json={"question": "hi"}).json()
    assert alice.get(f"/api/answers/{plain['answer_id']}/export.xlsx").status_code == 404


def test_question_to_a_deleted_chat_explains(web, monkeypatch):
    module, _ = web
    alice = login(module, "alice", "password-a")
    chat_id = alice.post("/api/chat", json={"question": "first"}).json()["conversation_id"]
    real_resolve = module._resolve

    def resolve_then_delete(user, request):
        result = real_resolve(user, request)
        module.store.delete_conversation(user["id"], chat_id)  # deleted in another tab meanwhile
        return result

    monkeypatch.setattr(module, "_resolve", resolve_then_delete)
    events = read_events(alice.post("/api/chat/stream", json={"question": "more", "conversation_id": chat_id}))
    assert events == [{"type": "error", "message": "Chat not found. It may have been deleted."}]


def test_old_failed_logins_are_forgotten(web):
    module, _ = web
    old = module.time.time() - module.LOCKOUT_SECONDS - 60
    module._failed_logins.update({f"typo{i}": [old] for i in range(600)})
    TestClient(module.app).post("/api/login", json={"username": "alice", "password": "wrong"})
    assert set(module._failed_logins) == {"alice"}


def test_chat_deleted_while_answering_explains(web, monkeypatch):
    module, _ = web
    alice = login(module, "alice", "password-a")
    chat_id = alice.post("/api/chat", json={"question": "first"}).json()["conversation_id"]
    user_id = module.store.authenticate("alice", "password-a")["id"]

    def delete_midway(history, question, model=None):
        module.store.delete_conversation(user_id, chat_id)  # deleted in another tab while answering
        yield {"type": "done", "answer": "late", "usage": None, "history": history}

    monkeypatch.setattr(module.agent, "respond_stream", delete_midway)
    events = read_events(alice.post("/api/chat/stream", json={"question": "more", "conversation_id": chat_id}))
    assert events[-1]["type"] == "error" and "deleted while David was answering" in events[-1]["message"]



def test_feedback_on_answers(web):
    module, _ = web
    alice = login(module, "alice", "password-a")  # admin
    bob = login(module, "bob", "password-b")
    asked = bob.post("/api/chat", json={"question": "Printer issues at Jimmy's?"}).json()
    answer_id, chat_id = asked["answer_id"], asked["conversation_id"]

    # Only the person who asked can rate their answer.
    assert alice.post(f"/api/answers/{answer_id}/feedback", json={"rating": 1}).status_code == 404
    assert bob.post(f"/api/answers/{answer_id}/feedback", json={"rating": 5}).status_code == 400
    assert bob.post(f"/api/answers/{answer_id}/feedback", json={"rating": -1}).status_code == 200
    assert bob.post(f"/api/answers/{answer_id}/feedback",
                    json={"rating": -1, "comment": "  Missed two tickets.  "}).status_code == 200

    # The rating shows again when the chat is reopened.
    turns = bob.get(f"/api/conversations/{chat_id}").json()["turns"]
    assert turns[1]["feedback"] == {"rating": -1, "comment": "Missed two tickets."} and turns[0]["feedback"] is None

    # Admins see everyone's feedback; others can't.
    assert bob.get("/api/feedback").status_code == 403
    assert bob.get("/feedback", follow_redirects=False).headers["location"] == "/"
    assert alice.get("/feedback").status_code == 200
    data = alice.get("/api/feedback").json()
    item = data["items"][0]
    assert data["down"] == 1 and data["up"] == 0
    assert item["display_name"] == "Bob B" and item["question"] == "Printer issues at Jimmy's?"
    assert item["comment"] == "Missed two tickets." and item["answer"].startswith("answer to")
    assert alice.get("/api/feedback/recent").json() == {"down_this_week": 1}  # just the count, for the icon
    assert bob.get("/api/feedback/recent").status_code == 403

    # Feedback stays readable after the chat is deleted, and can be taken back.
    bob.delete(f"/api/conversations/{chat_id}")
    assert alice.get("/api/feedback").json()["items"][0]["comment"] == "Missed two tickets."
    second = bob.post("/api/chat", json={"question": "hours?"}).json()["answer_id"]
    bob.post(f"/api/answers/{second}/feedback", json={"rating": 1})
    bob.post(f"/api/answers/{second}/feedback", json={"rating": 0})
    assert alice.get("/api/feedback").json()["up"] == 0


def test_sign_in_page_skips_ahead_when_signed_in(web):
    module, _ = web
    assert module.app and TestClient(module.app).get("/login").status_code == 200
    alice = login(module, "alice", "password-a")
    assert alice.get("/login", follow_redirects=False).headers["location"] == "/"


def test_overlong_questions_are_refused_before_reaching_claude(web, monkeypatch):
    module, _ = web
    client = login(module, "bob", "password-b")
    asked = []
    monkeypatch.setattr(module.agent, "respond_stream", lambda *a, **k: asked.append(a) or iter(()))
    for path in ("/api/chat", "/api/chat/stream"):
        too_long = client.post(path, json={"question": "x" * (module.MAX_QUESTION + 1)})
        assert too_long.status_code == 400 and "too long" in too_long.json()["detail"]
        assert client.post(path, json={"question": "   "}).status_code == 400
    assert asked == []


def spoton_zip() -> bytes:
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("about.txt", "Restaurant: Test Pub\nExported: 2026-10-06 10:00\n")
        z.writestr("menu_items.csv", "﻿Name,Price,ReportGroupName\nBurger,12.50,Food\nFries,4.00,\n")
    return buf.getvalue()


def test_only_admins_and_uploaders_can_upload_spoton_data(web):
    module, _ = web
    admin, user, uploader = (login(module, "alice", "password-a"), login(module, "bob", "password-b"),
                             login(module, "carol", "password-c"))
    assert admin.get("/api/me").json()["can_upload"] is True
    assert uploader.get("/api/me").json() | {"ticket_url": None} == {
        "username": "carol", "display_name": "Carol C", "is_admin": False, "can_upload": True, "ticket_url": None}
    assert user.get("/api/me").json()["can_upload"] is False

    assert user.post("/api/spoton/upload?filename=x.zip", content=spoton_zip()).status_code == 403
    assert user.get("/api/spoton").status_code == 403
    assert uploader.get("/api/feedback").status_code == 403  # uploaders aren't admins

    response = uploader.post("/api/spoton/upload?filename=Test_Pub.zip", content=spoton_zip())
    assert response.status_code == 200, response.text
    assert response.json() == {"restaurant": "Test Pub", "files": {"menu_items": 2}}
    assert [(f["restaurant"], f["file"], f["row_count"], f["uploaded_by"]) for f in admin.get("/api/spoton").json()] \
        == [("Test Pub", "menu_items", 2, "Carol C")]

    named = uploader.post("/api/spoton/upload?filename=Test Pub burgers.csv", content=b"Name\nBurger\n")
    assert named.json()["restaurant"] == "Test Pub"  # no question: the file name says which restaurant
    assert uploader.post("/api/spoton/upload?filename=menu.csv&restaurant=Other", content=b"Name\nBurger\n").status_code == 200

    assert user.delete("/api/spoton/Test Pub").status_code == 403
    assert admin.delete("/api/spoton/Test Pub").status_code == 200
    assert [f["restaurant"] for f in admin.get("/api/spoton").json()] == ["Other"]
    assert user.get("/api/spoton").status_code == 403


def test_a_question_about_an_upload_keeps_a_clean_title(web):
    module, calls = web
    alice = login(module, "alice", "password-a")
    question = "[SpotOn upload: Taco Town · Menu Items.xlsx]\nWhich items have no report group?"
    done = alice.post("/api/chat", json={"question": question}).json()
    assert done["title"] == "Which items have no report group?"
    assert calls[-1][1] == question  # David still sees which upload it's about


def test_single_files_add_to_a_restaurant_and_a_zip_replaces_it(web):
    module, _ = web
    admin = login(module, "alice", "password-a")
    files = lambda: sorted((f["restaurant"], f["file"], f["row_count"]) for f in admin.get("/api/spoton").json())  # noqa: E731
    up = lambda name, body, rest="": admin.post(f"/api/spoton/upload?filename={name}&restaurant={rest}", content=body)  # noqa: E731
    up("Menu Items.csv", b"Name\nBurger\nFries\n", "Taco Town")
    up("Employees.csv", b"First\nSam\n", "taco town")  # same restaurant, other spelling
    assert files() == [("Taco Town", "employees", 1), ("Taco Town", "menu_items", 2)]  # the first file stayed
    up("Menu Items.csv", b"Name\nBurger\n", "Taco Town")  # same file again: replaced
    assert files() == [("Taco Town", "employees", 1), ("Taco Town", "menu_items", 1)]
    assert admin.delete("/api/spoton/Taco Town?file=employees").status_code == 200
    assert admin.delete("/api/spoton/Taco Town?file=employees").status_code == 404
    assert files() == [("Taco Town", "menu_items", 1)]
    up("Test_Pub.zip", spoton_zip())
    up("Wings.csv", b"Name\nWings\n", "Test Pub")
    up("Test_Pub.zip", spoton_zip())  # a full export replaces everything for its restaurant
    assert files() == [("Taco Town", "menu_items", 1), ("Test Pub", "menu_items", 2)]

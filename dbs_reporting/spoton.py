"""SpotOn data uploaded from the SpotOn Exporter, and the tools David uses to read it.

The exporter's "Export CSVs" button makes one zip per restaurant: about.txt (the restaurant's name) plus
menu_items.csv, modifiers.csv, employees.csv and the Audit Check (audit_summary.csv, audit_items_to_fix.csv,
...). Admins and uploaders (see userfile.py) upload it; a new upload for a restaurant replaces the old one.
A single .csv or Excel workbook (.xlsx) can be uploaded too; it goes with the restaurant named in its file name
(see restaurant_for), or is listed under its own name. Each
non-empty sheet of a workbook becomes a file (named after the sheet, or after the workbook if it has one sheet).
"""

import csv
import datetime
import io
import json
import re
import zipfile
from collections import Counter

from anthropic import beta_tool

from .store import SPOTON_KEEP_DAYS, Store

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ROWS = 50_000  # per file
MAX_RESULT_ROWS = 1000
# Columns the exporter fills with comma-separated lists; group_by counts each entry on its own.
LIST_COLUMNS = {"MenuGroupNames", "Taxes", "ModifierGroups", "JobPositions", "Menu Groups"}
BLANK = "(blank)"


class UploadError(ValueError):
    """Something wrong with an uploaded file, worded for the person who uploaded it."""


def file_key(name: str) -> str:
    """"Audit Items To Fix.csv" -> "audit_items_to_fix"."""
    stem = re.sub(r"\.(csv|xlsx)$", "", name.rsplit("/", 1)[-1], flags=re.I)
    return re.sub(r"[^a-z0-9]+", "_", stem.lower()).strip("_") or "data"


def read_csv(data: bytes, name: str) -> tuple[list[str], list[list[str]]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return _table(csv.reader(io.StringIO(text, newline="")), name)


def _cell(value) -> str:
    """An Excel cell as the text a CSV would have: 12.5 -> "12.5", 3.0 -> "3", dates as 2026-10-06."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value).strip()


def read_xlsx(data: bytes, name: str) -> dict[str, tuple[list[str], list[list[str]]]]:
    """{file key: (columns, rows)} for each sheet with data. The first non-empty row of a sheet is its header."""
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UploadError(f"{name} can't be opened as an Excel workbook (.xlsx).")
    try:
        sheets = {}
        for sheet in book.worksheets:
            rows = ([_cell(v) for v in row] for row in sheet.iter_rows(values_only=True))
            rows = (row for row in rows if any(row))
            first = next(rows, None)
            if first is None:
                continue
            while first and not first[-1]:  # Excel often reports empty columns past the data
                first.pop()
            sheets[sheet.title] = _table(_chain([first], rows), f"{name} ({sheet.title})")
    finally:
        book.close()
    if not sheets:
        raise UploadError(f"{name} is empty.")
    if len(sheets) == 1:
        return {file_key(name): next(iter(sheets.values()))}
    return {file_key(title): table for title, table in sheets.items()}


def _chain(*iterables):
    for it in iterables:
        yield from it


def _table(reader, name: str) -> tuple[list[str], list[list[str]]]:
    """(columns, rows) from rows of text: the first row is the header; blank rows are skipped."""
    header = next(reader, None)
    if not header or not any(h.strip() for h in header):
        raise UploadError(f"{name} is empty.")
    columns = [h.strip() or f"Column {i + 1}" for i, h in enumerate(header)]
    rows = []
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        rows.append((row + [""] * len(columns))[:len(columns)])
        if len(rows) > MAX_ROWS:
            raise UploadError(f"{name} has more than {MAX_ROWS:,} rows.")
    return columns, rows


# What the SpotOn Exporter's files are called; at the end of a file name, what comes before is the restaurant.
SPOTON_FILE = re.compile(r"(?:^|[\s_\-–.]+)(menu[\s_]*items|menu[\s_]*groups|menus?|items|modifier[\s_]*groups|modifiers|"
                         r"employees|staff|exports?|spoton[\s_]*data|audit[\s_]*items[\s_]*to[\s_]*fix|audit[\s_]*summary|audit[\s_]*check|audit)$", re.I)


def _stem(filename: str) -> str:
    return re.sub(r"\.[a-z0-9]+$", "", filename.rsplit("/", 1)[-1], flags=re.I).strip()


def restaurant_for(filename: str, known: list[str]) -> str:
    """Which restaurant a single uploaded file is for, without asking: a known restaurant named in the file name
    ("Taco Town - Menu Items.xlsx" -> "Taco Town", the longest match wins); else what comes before a SpotOn file
    name ("Dock Bar employees.csv" -> "Dock Bar"); else the file's own name."""
    stem = _stem(filename)
    words = lambda text: " ".join(re.findall(r"[a-z0-9']+", text.casefold()))  # noqa: E731
    named = [k for k in known if words(k) and f" {words(k)} " in f" {words(stem)} "]
    if named:
        return max(named, key=len)
    found = SPOTON_FILE.search(stem)
    before = re.sub(r"[_\s]+", " ", stem[:found.start()]).strip(" -–_.") if found else ""
    return before or re.sub(r"[_\s]+", " ", stem).strip() or "Upload"


def single_file_key(filename: str, restaurant: str) -> str:
    """The file's name without the restaurant's: "Taco Town - Menu Items.xlsx" for Taco Town -> "menu_items"."""
    stem = _stem(filename)
    words = re.findall(r"[a-z0-9']+", stem.casefold())
    lead = re.findall(r"[a-z0-9']+", restaurant.casefold())
    rest = words[len(lead):] if lead and words[:len(lead)] == lead else words
    return file_key(" ".join(rest) or stem)


def parse_upload(filename: str, data: bytes, restaurant: str = "") -> tuple[str, dict]:
    """Returns (restaurant, {file key: (columns, rows)}) from an exporter zip or one .csv / .xlsx. A single file
    with no restaurant given is listed under its own name."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError(f"That file is over {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    restaurant = restaurant.strip()
    files = {}
    if filename.lower().endswith(".zip"):
        try:
            archive = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile:
            raise UploadError("That zip file can't be opened.")
        for info in archive.infolist():
            name = info.filename
            if info.is_dir() or name.rsplit("/", 1)[-1].startswith("."):
                continue
            if info.file_size > MAX_UPLOAD_BYTES:
                raise UploadError(f"{name} is too big.")
            if name.lower().endswith("about.txt") and not restaurant:
                about = archive.read(info).decode("utf-8-sig", "replace")
                found = re.search(r"^Restaurant:\s*(.+)$", about, re.M)
                restaurant = found.group(1).strip() if found else ""
            elif name.lower().endswith(".csv"):
                files[file_key(name)] = read_csv(archive.read(info), name)
            elif name.lower().endswith(".xlsx"):
                files.update(read_xlsx(archive.read(info), name))
        if not files:
            raise UploadError("No CSV files in that zip. Use the Export CSVs button in the SpotOn Exporter.")
    elif filename.lower().endswith(".csv"):
        files[single_file_key(filename, restaurant or restaurant_for(filename, []))] = read_csv(data, filename)
    elif filename.lower().endswith(".xlsx"):
        files = read_xlsx(data, filename)
        if len(files) == 1:  # one sheet: named after the file, without the restaurant's name
            files = {single_file_key(filename, restaurant or restaurant_for(filename, [])): next(iter(files.values()))}
    else:
        raise UploadError("Upload the .zip from the SpotOn Exporter's Export CSVs button, a .csv or an Excel .xlsx file.")
    if not restaurant:
        if filename.lower().endswith(".zip"):
            raise UploadError("That zip has no about.txt naming the restaurant. Use the SpotOn Exporter's Export CSVs button.")
        restaurant = restaurant_for(filename, [])
    return restaurant[:120], files


def _match(cell: str, wanted: str) -> bool:
    wanted = wanted.strip()
    if wanted == BLANK:
        return not cell.strip()
    if wanted.startswith("="):
        return cell.strip().lower() == wanted[1:].strip().lower()
    return wanted.lower() in cell.lower()


def build_spoton_tools(store: Store) -> list:
    def find_restaurant(name: str) -> tuple[str | None, list[str]]:
        names = sorted({f["restaurant"] for f in store.spoton_files()})
        exact = [n for n in names if n.lower() == name.strip().lower()]
        partial = [n for n in names if name.strip().lower() in n.lower()]
        found = exact or partial
        return (found[0] if len(found) == 1 else None), (found or names)

    @beta_tool(eager_input_streaming=True)
    def list_spoton_data() -> str:
        """List the SpotOn POS data people have uploaded: each restaurant, its files, row counts,
        columns, and who uploaded each file when.

        Call this first for any question about a restaurant's SpotOn menu, modifiers, employees or
        Audit Check, to see which restaurants and files exist. Uploads are snapshots taken for system
        audits and are deleted a month after upload, so say when the data you answer from was uploaded
        (e.g. "from the 10/6 upload").
        """
        try:
            restaurants: dict[str, dict] = {}
            for f in store.spoton_files():
                r = restaurants.setdefault(f["restaurant"], {"restaurant": f["restaurant"], "uploaded_by": f["uploaded_by"],
                                                             "uploaded_at": f["uploaded_at"], "files": []})
                r["uploaded_at"] = max(r["uploaded_at"], f["uploaded_at"])
                r["files"].append({"file": f["file"], "rows": f["row_count"], "uploaded_at": f["uploaded_at"],
                                   "columns": f["columns"]})
            if not restaurants:
                return json.dumps({"restaurants": [], "note": "No SpotOn data has been uploaded yet. An admin or "
                                   "uploader can upload it with the paperclip in the question box."})
            return json.dumps({"restaurants": list(restaurants.values()),
                               "note": f"Each file is deleted {SPOTON_KEEP_DAYS} days after it was uploaded."},
                              separators=(",", ":"), ensure_ascii=False)
        except Exception as exc:
            return json.dumps({"error": str(exc)})

    @beta_tool(eager_input_streaming=True)
    def get_spoton_data(restaurant: str, file: str, search: str = "", where: dict[str, str] | None = None,
                        columns: list[str] | None = None, group_by: str = "", limit: int = 300) -> str:
        """Read rows from one uploaded SpotOn file, optionally filtered, or count them by a column.

        Args:
            restaurant: Restaurant name as listed by list_spoton_data (part of the name is fine).
            file: File name from list_spoton_data, e.g. "menu_items", "modifiers", "employees", "audit_summary", "audit_items_to_fix".
            search: Optional text to find in any column (case-insensitive).
            where: Optional filters, {column: text}. Text matches if the cell contains it; "=Food" matches exactly; "(blank)" finds empty cells. All filters must match.
            columns: Optional columns to return (default all). Ask for only the ones you need on big files.
            group_by: Optional column to count rows by instead of listing them. List columns (MenuGroupNames, Taxes, ModifierGroups, JobPositions) count each entry.
            limit: Most rows to return (max 1000). The result says how many matched in total.
        """
        try:
            name, options = find_restaurant(restaurant)
            if name is None:
                return json.dumps({"error": f"No single uploaded restaurant matches {restaurant!r}.",
                                   "restaurants": options})
            saved = store.spoton_rows(name, file_key(file))
            if saved is None:
                have = [f["file"] for f in store.spoton_files() if f["restaurant"] == name]
                return json.dumps({"error": f"{name} has no {file!r} file.", "files": have})
            all_columns, rows = saved
            index = {c.lower(): i for i, c in enumerate(all_columns)}

            def col(c: str) -> int:
                if c.lower() not in index:
                    raise ValueError(f"No column {c!r}. Columns: {', '.join(all_columns)}")
                return index[c.lower()]

            filters = [(col(c), v) for c, v in (where or {}).items()]
            matched = [r for r in rows
                       if (not search or any(search.lower() in cell.lower() for cell in r))
                       and all(_match(r[i], v) for i, v in filters)]
            uploaded = next((f["uploaded_at"] for f in store.spoton_files()
                             if f["restaurant"] == name and f["file"] == file_key(file)), None)
            result = {"restaurant": name, "file": file_key(file), "uploaded_at": uploaded,
                      "total_rows": len(rows), "matched": len(matched)}
            if group_by:
                i = col(group_by)
                counts = Counter()
                for r in matched:
                    values = [v.strip() for v in r[i].split(",")] if all_columns[i] in LIST_COLUMNS else [r[i].strip()]
                    for v in values or [""]:
                        counts[v or BLANK] += 1
                result["group_by"] = all_columns[i]
                result["counts"] = counts.most_common(200)
                return json.dumps(result, separators=(",", ":"), ensure_ascii=False)
            keep = [col(c) for c in columns] if columns else list(range(len(all_columns)))
            limit = max(1, min(int(limit or 300), MAX_RESULT_ROWS))
            result["columns"] = [all_columns[i] for i in keep]
            result["rows"] = [[r[i] for i in keep] for r in matched[:limit]]
            if len(matched) > limit:
                result["note"] = f"Showing the first {limit} of {len(matched)} matching rows."
            return json.dumps(result, separators=(",", ":"), ensure_ascii=False)
        except Exception as exc:
            return json.dumps({"error": str(exc)})

    return [list_spoton_data, get_spoton_data]

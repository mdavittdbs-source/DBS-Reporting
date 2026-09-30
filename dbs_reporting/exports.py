"""Excel exports of the reports in an answer: each markdown table and each chart (with its data
and a native Excel chart). The answer's other text isn't included."""

import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Chart colors: the same validated categorical palette as the web page (light mode).
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
HEADER_FILL = PatternFill("solid", fgColor="E3E9F4")
NUMBER = re.compile(r"^-?\$?[\d,]*\.?\d+%?$")


def _plain(cell: str) -> str:
    """Strip markdown emphasis and links from a table cell."""
    cell = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), cell)
    cell = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*|`(.+?)`", lambda m: m.group(1) or m.group(2), cell)
    return re.sub(r"\[(.+?)\]\(.+?\)", r"\1", cell).strip()


def _value(cell: str):
    """Numbers become real numbers (so Excel can sum/sort them); everything else stays text."""
    text = _plain(cell)
    if NUMBER.match(text.replace(" ", "")):
        raw = text.replace(",", "").replace("$", "").replace(" ", "")
        try:
            return float(raw[:-1]) / 100 if raw.endswith("%") else float(raw) if "." in raw else int(raw)
        except ValueError:
            pass
    return text


def markdown_tables(text: str) -> list[list[list[str]]]:
    """Find pipe tables in markdown; each is a list of rows (header first) of raw cell text."""
    tables, current = [], []
    for line in text.splitlines() + [""]:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.count("|") >= 2:
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if len(current) == 1 and all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                continue  # the |---|---| separator row under the header
            current.append(cells)
        else:
            if len(current) >= 2:
                tables.append(current)
            current = []
    return tables


def _as_text(ws) -> None:
    """Keep text that starts with "=" as text. openpyxl would otherwise write it as a formula,
    which Excel then reports as a damaged file (or runs, if it came from a ticket summary)."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"


def _autosize(ws) -> None:
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(10, width + 2), 60)


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL


def _add_chart_sheet(wb: Workbook, number: int, chart: dict) -> None:
    ws = wb.create_sheet(f"Chart {number}")
    ws.append([chart["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    if chart.get("subtitle"):
        ws.append([chart["subtitle"]])
    header_row = ws.max_row + 2
    ws.cell(header_row, 1, chart.get("x_label") or "Category")
    for i, s in enumerate(chart["series"], start=2):
        ws.cell(header_row, i, s["name"])
    for r, label in enumerate(chart["labels"], start=header_row + 1):
        ws.cell(r, 1, label)
        for i, s in enumerate(chart["series"], start=2):
            ws.cell(r, i, s["values"][r - header_row - 1])
    _header(ws, header_row)
    _autosize(ws)

    last_row = header_row + len(chart["labels"])
    kind = chart["type"]
    xl = LineChart() if kind == "line" else BarChart()
    if kind != "line":
        xl.type = "bar" if kind == "hbar" else "col"
        xl.gapWidth = 60
        if kind == "stacked_bar":
            xl.grouping, xl.overlap = "stacked", 100
    xl.title = chart["title"]
    xl.y_axis.title = chart.get("y_label") or None
    xl.x_axis.title = chart.get("x_label") or None
    data = Reference(ws, min_col=2, max_col=1 + len(chart["series"]), min_row=header_row, max_row=last_row)
    xl.add_data(data, titles_from_data=True)
    xl.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
    for i, s in enumerate(xl.series):
        color = PALETTE[i % len(PALETTE)]
        if kind == "line":
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 25400  # 2pt
            s.smooth = False
        else:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.noFill = True
    if len(chart["series"]) == 1:
        xl.legend = None
    xl.height = max(7.5, 0.6 * len(chart["labels"]) + 3) if kind == "hbar" else 9
    xl.width = 22
    ws.add_chart(xl, f"{get_column_letter(len(chart['series']) + 3)}{header_row}")


def has_report(answer: dict) -> bool:
    """Whether an answer made anything to download: a chart or a table."""
    return bool(answer.get("charts")) or bool(markdown_tables(answer["text"]))


def answer_workbook(answer: dict) -> bytes:
    """Build the .xlsx for one answer (as returned by Store.get_answer): just its tables and
    charts, not the conversation. Check has_report() first."""
    wb = Workbook()
    wb.remove(wb.active)

    for number, table in enumerate(markdown_tables(answer["text"]), start=1):
        sheet = wb.create_sheet(f"Table {number}")
        for r, row in enumerate(table):
            sheet.append([_plain(c) if r == 0 else _value(c) for c in row])
        _header(sheet)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float) and table[cell.row - 1][cell.column - 1].strip().endswith("%"):
                    cell.number_format = "0.0%"
        _autosize(sheet)

    for number, chart in enumerate(answer.get("charts") or [], start=1):
        _add_chart_sheet(wb, number, chart)

    for sheet in wb.worksheets:
        _as_text(sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
